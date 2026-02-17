import frappe
from frappe.utils.pdf import get_pdf
from frappe.utils.xlsxutils import make_xlsx
from frappe import _

@frappe.whitelist()
def get_production_log_details(name=None):
    if not name:
        # If no name is provided, fetch the most recent log
        name = frappe.db.get_value("Daily Production Log", {"docstatus": ["<", 2]}, "name", order_by="creation desc")
    
    if not name:
        return None
        
    doc = frappe.get_doc("Daily Production Log", name)
    doc_dict = doc.as_dict()
    
    # child table data is already in doc.as_dict() for table fields
    # but let's ensure it's structured as the JS expects
    doc_dict["production_data"] = [d.as_dict() for d in doc.production_data]
    
    # Fetch RM UOM
    if doc.raw_material:
        doc_dict["rm_uom"] = frappe.db.get_value("Item", doc.raw_material, "stock_uom") or ""
    else:
        doc_dict["rm_uom"] = ""
    
    # Fetch Rejection Codes (Quality Inspection Parameters)
    params = frappe.get_all("Quality Inspection Parameter", fields=["name", "description"])
    doc_dict["rejection_codes_list"] = params
    
    return doc_dict

@frappe.whitelist()
def download_pdf(name):
    doc = get_production_log_details(name)
    if not doc:
        frappe.throw(_("Record not found"))
    
    html = f"""
    <html>
    <head>
        <style>
            body {{ font-family: sans-serif; font-size: 10px; margin: 0; padding: 5px; color: #000; }}
            .report-wrapper {{ border: 1px solid #000; padding: 10px; min-height: 900px; }}
            .report-table {{ width: 100%; border-collapse: collapse; border: 1px solid #000; table-layout: fixed; }}
            .report-table th, .report-table td {{ border: 1px solid #000; padding: 2px 5px; text-align: left; vertical-align: middle; height: 16px; }}
            .report-table th {{ background-color: #f2f2f2; text-align: center; text-transform: uppercase; font-size: 9px; font-weight: bold; }}
            .text-center {{ text-align: center !important; }}
            .font-bold {{ font-weight: bold; }}
            .label-cell {{ background-color: #f8f8f8; font-weight: bold; font-size: 9px; color: #000; border-right: none !important; text-transform: uppercase; }}
            .value-cell {{ font-weight: normal; font-size: 9px; color: #000; }}
            .title-large {{ font-size: 16px; font-weight: bold; }}
            .metadata-small {{ font-size: 8px; line-height: 1.1; color: #555; }}
            .rejection-codes {{ font-size: 9px; margin-top: 5px; border: 1px solid #000; padding: 5px; line-height: 1.3; }}
            .table-spacing {{ margin-top: -1.5px; }}
        </style>
    </head>
    <body>
        <div class="report-wrapper">
            <!-- Header -->
            <table class="report-table">
                <tr>
                    <td colspan="3" class="text-center font-bold" style="width: 30%;">{doc.get('company') or "YASH PLASTIC & ENGG WORKS"}</td>
                    <td colspan="4" class="text-center title-large" style="width: 40%;">DAILY PRODUCTION REPORT</td>
                    <td colspan="3" class="metadata-small" style="width: 30%;">
                        Doc No : {doc.get('doc_no') or ""},<br/>
                        REV No & Dt : {doc.get('rev_no') or ""}<br/>
                        Page : {doc.get('page_no') or "01"} of {doc.get('total_pages') or "01"}
                    </td>
                </tr>
                <tr>
                    <td colspan="2" class="label-cell">SHIFT DETAILS:</td>
                    <td colspan="2" class="text-center value-cell">{doc.get('shift') or ""}</td>
                    <td colspan="1" class="label-cell">DATE:</td>
                    <td colspan="2" class="text-center value-cell">{doc.get('report_date') or ""}</td>
                    <td colspan="2" class="label-cell">MACHINE NO:</td>
                    <td colspan="1" class="text-center value-cell">{doc.get('machine_no') or ""}</td>
                </tr>
                <tr>
                    <td colspan="2" class="label-cell">PRODUCT NAME:</td>
                    <td colspan="5" class="value-cell">{doc.get('product_name') or ""}</td>
                    <td colspan="3" class="value-cell"><span class="label-cell" style="background:transparent; border:none;">OPERATOR NAME:</span> <span style="margin-left: 5px;">{doc.get('operator_name') or ""}</span></td>
                </tr>
            </table>

            <!-- Metadata Grid -->
            <table class="report-table table-spacing">
                <tr>
                    <td colspan="2" class="label-cell" style="width: 20%;">SHOT WEIGHT:</td>
                    <td colspan="1" class="text-center value-cell" style="width: 10%;">{doc.get('shot_weight') or ""}</td>
                    <td colspan="2" class="label-cell" style="width: 20%;">RUNNER WEIGHT:</td>
                    <td colspan="1" class="text-center value-cell" style="width: 10%;">{doc.get('runner_weight') or ""}</td>
                    <td colspan="2" class="label-cell" style="width: 20%;">ITEM CODE NO:</td>
                    <td colspan="2" class="text-center value-cell" style="width: 20%;">{doc.get('item_code') or ""}</td>
                </tr>
                <tr>
                    <td colspan="2" class="label-cell">RAW MATERIAL:</td>
                    <td colspan="1" class="text-center value-cell">{doc.get('raw_material') or ""}</td>
                    <td colspan="2" class="label-cell">GRADE:</td>
                    <td colspan="1" class="text-center value-cell">{doc.get('raw_material_grade') or ""}</td>
                    <td colspan="2" class="label-cell">BATCH NO:</td>
                    <td colspan="2" class="text-center value-cell">{doc.get('raw_material_batch_no') or ""}</td>
                </tr>
                <tr>
                    <td colspan="2" class="label-cell">MASTERBATCH:</td>
                    <td colspan="1" class="text-center value-cell">{doc.get('masterbatch') or ""}</td>
                    <td colspan="2" class="label-cell">GRADE:</td>
                    <td colspan="1" class="text-center value-cell">{doc.get('masterbatch_grade') or ""}</td>
                    <td colspan="2" class="label-cell">BATCH NO:</td>
                    <td colspan="2" class="text-center value-cell">{doc.get('masterbatch_batch_no') or ""}</td>
                </tr>
                <tr>
                    <td colspan="2" class="label-cell">FIRST COUNTER:</td>
                    <td colspan="1" class="text-center value-cell">{doc.get('first_counter') or ""}</td>
                    <td colspan="2" class="label-cell">CYCLE TIME:</td>
                    <td colspan="1" class="text-center value-cell">{doc.get('cycle_time') or ""}</td>
                    <td colspan="2" class="label-cell">SHIFT TARGET:</td>
                    <td colspan="2" class="text-center value-cell">{doc.get('shift_target') or ""}</td>
                </tr>
            </table>

            <table class="report-table table-spacing">
                <tr>
                   <td colspan="2" class="label-cell" style="width: 20%;">TOTAL NO CAVITY:</td>
                   <td colspan="1" class="text-center value-cell" style="width: 10%;">{doc.get('total_cavity') or ""}</td>
                   <td colspan="2" class="label-cell" style="width: 20%;">RUNNING CAVITY:</td>
                   <td colspan="1" class="text-center value-cell" style="width: 10%;">{doc.get('running_cavity') or ""}</td>
                   <td colspan="2" class="label-cell" style="width: 20%;">ANTI STATIC:</td>
                   <td colspan="2" class="text-center value-cell" style="width: 20%;">{doc.get('anti_static') or ""}</td>
                </tr>
            </table>

            <table class="report-table table-spacing">
                <thead>
                    <tr>
                        <th style="width: 10%;">TIME</th>
                        <th style="width: 10%;">OK SHOTS</th>
                        <th style="width: 10%;">REJ SHOTS</th>
                        <th style="width: 10%;">TOTAL SHOTS</th>
                        <th style="width: 20%;">REJ CODE</th>
                        <th style="width: 40%;">REMARKS</th>
                    </tr>
                </thead>
                <tbody>
    """
    
    for row in doc.get("production_data", []):
        html += f"""
                <tr>
                    <td class="text-center value-cell">{row.get('time_slot') or ''}</td>
                    <td class="text-center value-cell">{row.get('ok_shots') or 0}</td>
                    <td class="text-center value-cell">{row.get('rej_shots') or 0}</td>
                    <td class="text-center value-cell">{row.get('total_shots') or 0}</td>
                    <td class="text-center value-cell">{row.get('rej_code') or ''}</td>
                    <td>{row.get('remarks') or ''}</td>
                </tr>
        """
        
    html += f"""
                </tbody>
            </table>
            
            <table class="report-table table-spacing">
                <tr>
                    <td class="label-cell text-center" style="width: 12%;">LAST COUNTER</td>
                    <td class="text-center value-cell" style="width: 10%; border-right: 2px solid #000;">{doc.get('last_counter') or ""}</td>
                    <td class="text-center value-cell" style="width: 10%;">{doc.get('total_ok_shots') or 0}</td>
                    <td class="text-center value-cell" style="width: 10%;">{doc.get('total_rej_shots') or 0}</td>
                    <td class="text-center value-cell" style="width: 10%; border-right: 2px solid #000;">{doc.get('total_shots') or 0}</td>
                    <td class="label-cell text-center" style="width: 18%;">RM CONSUMPTION</td>
                    <td class="label-cell text-center" style="width: 10%;">LUMPS</td>
                    <td class="label-cell text-center font-bold" style="width: 20%;">SHIFT SUPERVISOR/ QC SIGN</td>
                </tr>
                <tr>
                    <td colspan="5" style="border-right: 1.5px solid #000;"></td>
                    <td class="text-center value-cell" style="height: 45px;">{doc.get('rm_consumption') or 0} {doc.get('rm_uom', '')}</td>
                    <td class="text-center value-cell">{doc.get('lumps') or ""}</td>
                    <td style="vertical-align: bottom; height: 45px;">
                        <div style="text-align: center; font-style: italic; color: #777;">{ "Signed" if doc.get('supervisor_sign') else "" }</div>
                    </td>
                </tr>
            </table>

            <div class="rejection-codes">
                <strong style="font-size: 8px; color: #444;">DEFECT CODES:</strong><br/>
                {", ".join([f"{c['name']}{' - ' + c['description'] if c.get('description') else ''}" for c in doc.get('rejection_codes_list', [])])}
            </div>
        </div>
    </body>
    </html>
    """
    
    pdf_content = get_pdf(html)
    
    frappe.local.response.filename = f"{doc.get('name')}.pdf"
    frappe.local.response.filecontent = pdf_content
    frappe.local.response.type = "pdf"

@frappe.whitelist()
def download_excel(name):
    doc = get_production_log_details(name)
    if not doc:
        frappe.throw(_("Record not found"))
    
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Production Report"
    
    # Styles
    thin_side = Side(style='thin', color="000000")
    all_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    title_font = Font(name='Arial', size=14, bold=True)
    header_font = Font(name='Arial', size=11, bold=True)
    label_font = Font(name='Arial', size=9, bold=True)
    value_font = Font(name='Arial', size=9)
    meta_font = Font(name='Arial', size=8, color="555555")
    
    label_fill = PatternFill(start_color='F8F8F8', end_color='F8F8F8', fill_type='solid')
    header_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    # Set column widths (10 columns A-J)
    for col in range(1, 11):
        ws.column_dimensions[get_column_letter(col)].width = 14
    
    def apply_style(range_str, font=None, border=None, fill=None, align=None):
        for row_objs in ws[range_str]:
            for cell in row_objs:
                if font: cell.font = font
                if border: cell.border = border
                if fill: cell.fill = fill
                if align: cell.alignment = align

    # Row 1-2: Header
    ws.merge_cells('A1:C2')
    ws['A1'] = doc.get('company') or "YASH PLASTIC & ENGG WORKS"
    apply_style('A1:C2', font=header_font, border=all_border, align=center_align)
    
    ws.merge_cells('D1:G2')
    ws['D1'] = "DAILY PRODUCTION REPORT"
    apply_style('D1:G2', font=title_font, border=all_border, align=center_align)
    
    ws.merge_cells('H1:J2')
    meta_text = f"Doc No : {doc.get('doc_no') or ''}\nREV No & Dt : {doc.get('rev_no') or ''}\nPage : {doc.get('page_no') or '01'} of {doc.get('total_pages') or '01'}"
    ws['H1'] = meta_text
    apply_style('H1:J2', font=meta_font, border=all_border, align=left_align)
    
    # Row 3: Shift, Date, Machine
    ws.merge_cells('A3:B3')
    ws['A3'] = "SHIFT DETAILS:"
    ws.merge_cells('C3:D3')
    ws['C3'] = doc.get('shift') or ""
    ws['E3'] = "DATE:"
    ws.merge_cells('F3:G3')
    ws['F3'] = doc.get('report_date') or ""
    ws.merge_cells('H3:I3')
    ws['H3'] = "MACHINE NO:"
    ws['J3'] = doc.get('machine_no') or ""
    apply_style('A3:B3', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('C3:D3', font=value_font, border=all_border, align=center_align)
    apply_style('E3:E3', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('F3:G3', font=value_font, border=all_border, align=center_align)
    apply_style('H3:I3', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('J3:J3', font=value_font, border=all_border, align=center_align)
    
    # Row 4: Product Name, Operator
    ws.merge_cells('A4:B4')
    ws['A4'] = "PRODUCT NAME:"
    ws.merge_cells('C4:G4')
    ws['C4'] = doc.get('product_name') or ""
    ws.merge_cells('H4:I4')
    ws['H4'] = "OPERATOR NAME:"
    ws['J4'] = doc.get('operator_name') or ""
    apply_style('A4:B4', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('C4:G4', font=value_font, border=all_border, align=left_align)
    apply_style('H4:I4', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('J4:J4', font=value_font, border=all_border, align=center_align)
    
    # Row 5: Weights
    ws.merge_cells('A5:B5')
    ws['A5'] = "SHOT WEIGHT:"
    ws['C5'] = doc.get('shot_weight') or ""
    ws.merge_cells('D5:E5')
    ws['D5'] = "RUNNER WEIGHT:"
    ws['F5'] = doc.get('runner_weight') or ""
    ws.merge_cells('G5:H5')
    ws['G5'] = "ITEM CODE NO:"
    ws.merge_cells('I5:J5')
    ws['I5'] = doc.get('item_code') or ""
    apply_style('A5:B5', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('C5:C5', font=value_font, border=all_border, align=center_align)
    apply_style('D5:E5', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('F5:F5', font=value_font, border=all_border, align=center_align)
    apply_style('G5:H5', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('I5:J5', font=value_font, border=all_border, align=center_align)
    
    # Row 6: Material
    ws.merge_cells('A6:B6')
    ws['A6'] = "RAW MATERIAL:"
    ws['C6'] = doc.get('raw_material') or ""
    ws.merge_cells('D6:E6')
    ws['D6'] = "GRADE:"
    ws['F6'] = doc.get('raw_material_grade') or ""
    ws.merge_cells('G6:H6')
    ws['G6'] = "BATCH NO:"
    ws.merge_cells('I6:J6')
    ws['I6'] = doc.get('raw_material_batch_no') or ""
    apply_style('A6:B6', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('C6:C6', font=value_font, border=all_border, align=center_align)
    apply_style('D6:E6', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('F6:F6', font=value_font, border=all_border, align=center_align)
    apply_style('G6:H6', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('I6:J6', font=value_font, border=all_border, align=center_align)
    
    # Row 7: Masterbatch
    ws.merge_cells('A7:B7')
    ws['A7'] = "MASTERBATCH:"
    ws['C7'] = doc.get('masterbatch') or ""
    ws.merge_cells('D7:E7')
    ws['D7'] = "GRADE:"
    ws['F7'] = doc.get('masterbatch_grade') or ""
    ws.merge_cells('G7:H7')
    ws['G7'] = "BATCH NO:"
    ws.merge_cells('I7:J7')
    ws['I7'] = doc.get('masterbatch_batch_no') or ""
    apply_style('A7:B7', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('C7:C7', font=value_font, border=all_border, align=center_align)
    apply_style('D7:E7', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('F7:F7', font=value_font, border=all_border, align=center_align)
    apply_style('G7:H7', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('I7:J7', font=value_font, border=all_border, align=center_align)
    
    # Row 8: Counters
    ws.merge_cells('A8:B8')
    ws['A8'] = "FIRST COUNTER:"
    ws['C8'] = doc.get('first_counter') or ""
    ws.merge_cells('D8:E8')
    ws['D8'] = "CYCLE TIME:"
    ws['F8'] = doc.get('cycle_time') or ""
    ws.merge_cells('G8:H8')
    ws['G8'] = "SHIFT TARGET:"
    ws.merge_cells('I8:J8')
    ws['I8'] = doc.get('shift_target') or ""
    apply_style('A8:B8', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('C8:C8', font=value_font, border=all_border, align=center_align)
    apply_style('D8:E8', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('F8:F8', font=value_font, border=all_border, align=center_align)
    apply_style('G8:H8', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('I8:J8', font=value_font, border=all_border, align=center_align)
    
    # Row 9: Cavity
    ws.merge_cells('A9:B9')
    ws['A9'] = "TOTAL NO CAVITY:"
    ws['C9'] = doc.get('total_cavity') or ""
    ws.merge_cells('D9:E9')
    ws['D9'] = "RUNNING CAVITY:"
    ws['F9'] = doc.get('running_cavity') or ""
    ws.merge_cells('G9:H9')
    ws['G9'] = "ANTI STATIC:"
    ws.merge_cells('I9:J9')
    ws['I9'] = doc.get('anti_static') or ""
    apply_style('A9:B9', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('C9:C9', font=value_font, border=all_border, align=center_align)
    apply_style('D9:E9', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('F9:F9', font=value_font, border=all_border, align=center_align)
    apply_style('G9:H9', font=label_font, fill=label_fill, border=all_border, align=left_align)
    apply_style('I9:J9', font=value_font, border=all_border, align=center_align)
    
    # Row 11: Table Header
    ws['A11'] = "TIME"
    ws['B11'] = "OK SHOTS"
    ws['C11'] = "REJ SHOTS"
    ws['D11'] = "TOTAL SHOTS"
    ws.merge_cells('E11:F11')
    ws['E11'] = "REJ CODE"
    ws.merge_cells('G11:J11')
    ws['G11'] = "REMARKS"
    apply_style('A11:J11', font=label_font, fill=header_fill, border=all_border, align=center_align)
    
    curr_row = 12
    for row_data in doc.get("production_data", []):
        ws.cell(row=curr_row, column=1, value=row_data.get("time_slot") or "")
        ws.cell(row=curr_row, column=2, value=row_data.get("ok_shots") or 0)
        ws.cell(row=curr_row, column=3, value=row_data.get("rej_shots") or 0)
        ws.cell(row=curr_row, column=4, value=row_data.get("total_shots") or 0)
        ws.merge_cells(start_row=curr_row, start_column=5, end_row=curr_row, end_column=6)
        ws.cell(row=curr_row, column=5).value = row_data.get("rej_code") or ""
        ws.merge_cells(start_row=curr_row, start_column=7, end_row=curr_row, end_column=10)
        ws.cell(row=curr_row, column=7).value = row_data.get("remarks") or ""
        
        apply_style(f'A{curr_row}:J{curr_row}', font=value_font, border=all_border, align=center_align)
        curr_row += 1
    
    # Summary Footer
    ws['A'+str(curr_row)] = "LAST COUNTER"
    ws['B'+str(curr_row)] = doc.get("last_counter") or ""
    ws['C'+str(curr_row)] = doc.get("total_ok_shots") or 0
    ws['D'+str(curr_row)] = doc.get("total_rej_shots") or 0
    ws['E'+str(curr_row)] = doc.get("total_shots") or 0
    ws.merge_cells(f'F{curr_row}:G{curr_row}')
    ws['F'+str(curr_row)] = "RM CONSUMPTION"
    ws['H'+str(curr_row)] = "LUMPS"
    ws.merge_cells(f'I{curr_row}:J{curr_row}')
    ws['I'+str(curr_row)] = "SHIFT SUPERVISOR/ QC SIGN"
    
    apply_style(f'A{curr_row}:A{curr_row}', font=label_font, fill=label_fill, border=all_border, align=center_align)
    apply_style(f'B{curr_row}:E{curr_row}', font=value_font, border=all_border, align=center_align)
    apply_style(f'F{curr_row}:H{curr_row}', font=label_font, fill=label_fill, border=all_border, align=center_align)
    apply_style(f'I{curr_row}:J{curr_row}', font=label_font, fill=label_fill, border=all_border, align=center_align)
    
    curr_row += 1
    ws.merge_cells(f'A{curr_row}:E{curr_row}')
    ws.merge_cells(f'F{curr_row}:G{curr_row}')
    ws['F'+str(curr_row)] = f"{doc.get('rm_consumption') or 0} {doc.get('rm_uom', '')}"
    ws['H'+str(curr_row)] = doc.get("lumps") or ""
    ws.merge_cells(f'I{curr_row}:J{curr_row}')
    ws['I'+str(curr_row)] = "Signed" if doc.get('supervisor_sign') else ""
    
    apply_style(f'A{curr_row}:J{curr_row}', font=value_font, border=all_border, align=center_align)
    
    # Rejection Codes
    curr_row += 2
    ws.merge_cells(f'A{curr_row}:J{curr_row+1}')
    rej_codes_text = "DEFECT CODES: " + ", ".join([f"{c['name']}{' - ' + c['description'] if c.get('description') else ''}" for c in doc.get('rejection_codes_list', [])])
    ws['A'+str(curr_row)] = rej_codes_text
    apply_style(f'A{curr_row}:J{curr_row+1}', font=value_font, border=all_border, align=left_align)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    frappe.local.response.filename = f"{doc.get('name')}.xlsx"
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type = "binary"
